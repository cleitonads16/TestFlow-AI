import pytest
from docx import Document

from documentos import (
    DocumentoIlegivel,
    DocumentoSemTexto,
    FormatoDocumentoNaoSuportado,
    extrair_texto,
    validar_documento,
)


def test_extrai_texto_de_docx(tmp_path):
    caminho = tmp_path / "escopo.docx"
    documento = Document()
    documento.add_paragraph("Título do Escopo")
    documento.add_paragraph("Este é o corpo do escopo de teste.")
    documento.save(caminho)

    texto = extrair_texto(caminho)

    assert "Título do Escopo" in texto
    assert "corpo do escopo de teste" in texto


def test_extrai_tabelas_na_ordem_do_documento(tmp_path):
    caminho = tmp_path / "escopo.docx"
    documento = Document()
    documento.add_paragraph("Requisitos funcionais")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "RF01"
    tabela.cell(0, 1).text = "Login com SSO"
    tabela.cell(1, 0).text = "RF02"
    tabela.cell(1, 1).text = "Recuperar senha"
    documento.add_paragraph("Fim do escopo")
    documento.save(caminho)

    texto = extrair_texto(caminho)

    assert texto.splitlines() == [
        "Requisitos funcionais",
        "RF01 | Login com SSO",
        "RF02 | Recuperar senha",
        "Fim do escopo",
    ]


def test_celula_mesclada_aparece_uma_vez_so(tmp_path):
    caminho = tmp_path / "escopo.docx"
    documento = Document()
    tabela = documento.add_table(rows=1, cols=3)
    tabela.cell(0, 0).merge(tabela.cell(0, 1)).text = "Cadastro de clientes"
    tabela.cell(0, 2).text = "Alta"
    documento.save(caminho)

    assert extrair_texto(caminho) == "Cadastro de clientes | Alta"


def test_documento_sem_texto_levanta_documento_sem_texto(tmp_path):
    caminho = tmp_path / "escopo.docx"
    documento = Document()
    documento.add_paragraph("   ")
    documento.add_table(rows=1, cols=2)
    documento.save(caminho)

    with pytest.raises(DocumentoSemTexto, match="não tem texto"):
        extrair_texto(caminho)


@pytest.mark.parametrize("nome", ["escopo.txt", "escopo.pdf"], ids=["txt", "pdf"])
def test_formato_nao_suportado_levanta_erro(tmp_path, nome):
    caminho = tmp_path / nome
    caminho.write_text("conteúdo qualquer", encoding="utf-8")

    with pytest.raises(FormatoDocumentoNaoSuportado, match=r"Use \.docx"):
        extrair_texto(caminho)


def test_arquivo_ilegivel_levanta_documento_ilegivel(tmp_path):
    caminho = tmp_path / "escopo.docx"
    caminho.write_bytes(b"texto qualquer, so renomeado")

    with pytest.raises(DocumentoIlegivel, match="corrompido"):
        extrair_texto(caminho)


def test_validar_documento_aceita_docx_valido(conteudo_docx):
    validar_documento(conteudo_docx, ".docx")


def test_validar_documento_recusa_conteudo_ilegivel():
    with pytest.raises(DocumentoIlegivel):
        validar_documento(b"nao sou um docx", ".docx")
